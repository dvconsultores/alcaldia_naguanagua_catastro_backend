from rest_framework import viewsets
from rest_framework.views import APIView

from rest_framework.permissions import IsAuthenticated, AllowAny
from .serializers import *
from .use_cases import *
from django.contrib.auth.models import *
from rest_framework.decorators import api_view,permission_classes
from django.views.decorators.csrf import csrf_exempt
from django_filters.rest_framework import DjangoFilterBackend
from django_filters import rest_framework as filters
from django.db.models import Q
from django.http import JsonResponse
import json
from django.db.models import Sum, Prefetch


import pandas as pd
from io import BytesIO
from datetime import datetime
from django.core.files.base import ContentFile



import os
import openpyxl
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_POST
from django.shortcuts import render
from .models import ExcelDocument

@csrf_exempt
@require_POST
def subir_archivo_excel(request):
    archivo_excel = request.FILES['archivoExcel']

    # Verifica si el archivo es un archivo Excel válido
    if archivo_excel.name.endswith(('.xls', '.xlsx')):
        # Define la ubicación donde deseas guardar el archivo Excel
        ubicacion_archivo = os.path.join('media', 'archivos_excel', archivo_excel.name)

        # Guarda el archivo Excel en la ubicación especificada
        with open(ubicacion_archivo, 'wb') as archivo_destino:
            for chunk in archivo_excel.chunks():
                archivo_destino.write(chunk)

        # Ejemplo de procesamiento del archivo Excel (puede variar según tus necesidades)
        workbook = openpyxl.load_workbook(ubicacion_archivo)
        # Procesa el archivo Excel aquí, por ejemplo, lee hojas de trabajo y realiza operaciones

        return JsonResponse({'mensaje': 'Archivo Excel subido y procesado correctamente.'})
    else:
        return JsonResponse({'error': 'Formato de archivo no válido.'}, status=400)

@csrf_exempt
@require_POST
def upload_excel(request):
    if request.method == 'POST':
        title = request.POST['title']
        excel_file = request.FILES['excel_file']
        ExcelDocument.objects.filter(title=title,).delete()
        document = ExcelDocument(title=title, excel_file=excel_file)
        document.save()
        return JsonResponse({'message': 'Excel file uploaded successfully','id':document.id})
    return JsonResponse({'message': 'No file uploaded'}, status=400)


class MultiSerializerViewSet(viewsets.ModelViewSet):
    serializers = { 
        'default': None,
    }

    def get_serializer_class(self):
            return self.serializers.get(self.action,
                        self.serializers.get('default'))


@api_view(["POST"])
@csrf_exempt
@permission_classes([AllowAny])
def SignIn(request):
    return generate_token(request.data["username"],request.data["password"])

@api_view(["POST"])
@csrf_exempt
@permission_classes([AllowAny])
def SignUp(request):
    return create_user(request.data["username"],request.data["password"],request.data["email"])

@api_view(["POST"])
@csrf_exempt
@permission_classes([IsAuthenticated])
def ChangePassword(request):
    return change_password(request.user,request.data["password"])


@api_view(["POST"])
@csrf_exempt
@permission_classes([IsAuthenticated])
def CrearPatente(request):
    datos=request.data
    return Crear_Patente(datos)

@api_view(["POST"])
@csrf_exempt
@permission_classes([IsAuthenticated])
def CrearEstadoCuenta(request):
    datos=request.data
    return Crear_Estado_Cuenta(datos)

@api_view(["POST"])
@csrf_exempt
@permission_classes([IsAuthenticated])
def CrearLiquidacion(request):
    datos=request.data
    return Crear_Liquidacion(datos)

@api_view(["POST"])
@csrf_exempt
@permission_classes([IsAuthenticated])
def CrearPago(request):
    datos=request.data
    return Crear_Pago(datos)

@api_view(["POST"])
@csrf_exempt
@permission_classes([IsAuthenticated])
def CrearInmueblePropietario(request):
    datos=request.data
    return Crear_Inmueble_Propietario(datos)

@api_view(["POST"])
@csrf_exempt
@permission_classes([IsAuthenticated])
def MultaInmueble(request):
    datos=request.data
    return Multa_Inmueble(datos)


@api_view(["POST"])
@csrf_exempt
@permission_classes([IsAuthenticated])
def CrearPerfl(request):
    datos=request.data
    return Crear_Perfil(datos)

@api_view(["POST"])
@csrf_exempt
@permission_classes([IsAuthenticated])
def ImpuestoInmueble(request):
    datos=request.data
    print('datos',datos)
    return Impuesto_Inmueble(datos)

@api_view(["POST"])
@csrf_exempt
@permission_classes([IsAuthenticated])
def ImpuestoInmueble2023(request):
    datos=request.data
    print('datos2023',datos)
    return Impuesto_Inmueble2023(datos)

@api_view(["POST"])
@csrf_exempt
@permission_classes([AllowAny])
def EstadisticaFlujo(request):
    datos=request.data
    return Estadistica_Flujo(datos)


# API PUBLICA!!!!!!!!!!!!
@api_view(["POST"])
@csrf_exempt
@permission_classes([AllowAny])
def ImpuestoInmueblePublic(request):
    datos=request.data
    return Impuesto_Inmueble_Public(datos)

# API PUBLICA!!!!!!!!!!!!
@api_view(["POST"])
@csrf_exempt
@permission_classes([AllowAny])
def ImpuestoInmueble2023Public(request):
    datos=request.data
    return Impuesto_Inmueble2023_Public(datos)

# API PUBLICA!!!!!!!!!!!!
@api_view(["POST"])
@csrf_exempt
@permission_classes([AllowAny])
def ImpuestoInmueblePago(request):
    datos=request.data
    return Impuesto_Inmueble_Pago(datos)  

@api_view(["POST"])
@csrf_exempt
@permission_classes([AllowAny])
def DatosInmueblesPublic(request):
    datos=request.data
    return Datos_Inmuebles_Public(datos)
  
# API PUBLICA!!!!!!!!!!!!
@api_view(["GET"])
@csrf_exempt
@permission_classes([AllowAny])
def CertificaFicha(request):
    datos = request.query_params
    data_string = datos.get('data')  # Obtenemos la cadena JSON
    data_json = json.loads(data_string)  # Convertimos la cadena JSON a un diccionario
    ficha = data_json.get('ficha')  # Obtenemos el valor de "ficha" del diccionario
    print('CertificaFicha', ficha)
    return Certifica_Ficha(ficha)



@api_view(["POST"])
@csrf_exempt
@permission_classes([IsAuthenticated])
def ImpuestoInmuebleDetalle(request):
    datos=request.data
    return Impuesto_Inmueble_Detalle(datos)



###################################################

@api_view(["POST"])
@csrf_exempt
@permission_classes([IsAuthenticated])
def importardatosdesdeexcel(request):
    datos=request.data
    return importar_datos_desde_excel(datos['archivoExcel'],datos['opcion'])


@api_view(["POST"])
@csrf_exempt
@permission_classes([IsAuthenticated])
def importarcorridabancaria(request):
    datos=request.data
    return importar_corrida_bancaria(datos['archivoExcel'],datos['opcion'],datos['ruta'])

###################################################

class UserViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=User.objects.all()
    serializers = {
        'default': UserSerializer
    }




class DepartamentoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset = Departamento.objects.all()
    serializers = {
        'default': DepartamentoSerializer
    }
    filter_backends = [DjangoFilterBackend]
# excluye el departamento de usuario actual.
# este filtro se usa cuando enviamos un documento en FLUJO. y con esto
# evitamos que se envie el doucmento al mismo departamento que envia.
    def get_queryset(self):
        queryset = super().get_queryset()
        nombre_a_excluir = self.request.query_params.get('nombre')
        if nombre_a_excluir:
            queryset = queryset.exclude(nombre=nombre_a_excluir)
        return queryset



class PerfilViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Perfil.objects.all()
    serializers = {
        'default': PerfilSerializer
    }

class ModuloViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Modulo.objects.all()
    serializers = {
        'default': ModuloSerializer
    }

class PermisoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Permiso.objects.all()
    serializers = {
        'default': PermisoSerializer
    }

class AmbitoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Ambito.objects.all()
    serializers = {
        'default': AmbitoSerializer,
        'create':CreateAmbitoSerializer
    }


class SectorViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Sector.objects.all()
    serializers = {
        'default': SectorSerializer
    }

class CalleViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Calle.objects.all()
    serializers = {
        'default': CalleSerializer
    }

class AvenidaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Avenida.objects.all()
    serializers = {
        'default': AvenidaSerializer
    }

class UrbanizacionViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Urbanizacion.objects.all()
    serializers = {
        'default': UrbanizacionSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'id':['exact'],
    }

class ManzanaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Manzana.objects.all()
    serializers = {
        'default': ManzanaSerializer
    }

class ParcelaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Parcela.objects.all()
    serializers = {
        'default': ParcelaSerializer
    }

class SubParcelaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=SubParcela.objects.all()
    serializers = {
        'default': SubParcelaSerializer
    }

class ConjuntoResidencialViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=ConjuntoResidencial.objects.all()
    serializers = {
        'default': ConjuntoResidencialSerializer
    }

class EdificioViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Edificio.objects.all()
    serializers = {
        'default': EdificioSerializer
    }

class TorreViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Torre.objects.all()
    serializers = {
        'default': TorreSerializer
    }


class PropietarioViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset = Propietario.objects.all()
    serializers = {
        'default': PropietarioSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'numero_documento':['exact'],
    }
#++++++++++++++++++++++++++++++++++++++++++++++++++++

@csrf_exempt
def filtrar_propietarios(request):
    numero_documento = request.GET.get('numero_documento', None)

    #if numero_documento:
    #propietarios_filtrados = Propietario.objects.filter(numero_documento__contains=numero_documento)
    propietarios_filtrados = Propietario.objects.filter(
    Q(numero_documento__icontains=numero_documento) | Q(nombre__icontains=numero_documento)
    )
    #else:
    #    propietarios_filtrados = Propietario.objects.all()

    # Convierte los resultados a JSON
    data = [{'id': prop.id, 
             'numero_documento': prop.numero_documento, 
             'nombre': prop.nombre,
             'tipo_documento': prop.tipo_documento,
             'nacionalidad': prop.nacionalidad,
             'telefono_principal': prop.telefono_principal,
             'telefono_secundario': prop.telefono_secundario,
             'email_principal': prop.email_principal,
             'emaill_secundario': prop.emaill_secundario,
             'direccion': prop.direccion
             } for prop in propietarios_filtrados]

    return JsonResponse(data, safe=False)


@csrf_exempt
def filtrar_inmuebles(request):
    numero_expediente = request.GET.get('numero_expediente', None)
    inmuebles_filtrados = Inmueble.objects.filter(numero_expediente=numero_expediente)
    data = [
        {'id': prop.id, 
            'numero_expediente': prop.numero_expediente,
            'fecha_inscripcion': prop.fecha_inscripcion,
            'numero_civico': prop.numero_civico,
            'numero_casa': prop.numero_casa,
            'numero_piso': prop.numero_piso,
            'telefono': prop.telefono,
            #'zona': prop.zona,
            'direccion': prop.direccion, 
            'referencia': prop.referencia,
            'observaciones': prop.observaciones,
            'inscripcion_paga': prop.inscripcion_paga,
            'habilitado': prop.habilitado,
            #'periodo': prop.periodo,
            'anio ': prop.anio
        } 
        for prop in inmuebles_filtrados
        ]

    return JsonResponse(data, safe=False)

@csrf_exempt
def listar_inmuebles_old(request):
    inmuebles_filtrados = Inmueble.objects.all()
    data = []
    today = date.today() # año vigente para evaluar si los pagos estan vigentes o no 
    for prop in inmuebles_filtrados:
        try:
            terreno2024 = InmuebleValoracionTerreno2024.objects.get(inmueble__numero_expediente=prop.numero_expediente)
            total_area_terreno2024=terreno2024.area or 0
            construccion2024 = InmuebleValoracionConstruccion2024.objects.filter(inmueblevaloracionterreno=terreno2024) 
            total_area_construccion2024 = construccion2024.aggregate(Sum('area'))['area__sum']
            total_area_construccion2024 = total_area_construccion2024 or 0
        except InmuebleValoracionTerreno2024.DoesNotExist:
            terreno2024 = None
            total_area_construccion2024 = None

                # Obtener y concatenar propietarios
        oInmueblePropietarios = InmueblePropietarios.objects.filter(inmueble=prop.id)
        propietarios = ", ".join([f"{propietario.propietario.nombre} ({propietario.propietario.numero_documento})" for propietario in oInmueblePropietarios])
        # Verificar si anio es None antes de la comparación
        if prop.anio is not None:
            estatus = "Vigente" if prop.anio >= today.year else "No Vigente"
        else:
            estatus = "No Vigente"
        print(prop.numero_expediente)
        data.append({
            'estatus': estatus,
            'expediente': prop.numero_expediente,
            'zona': None if prop.zona is None else prop.zona.codigo,
            'categorizacion': None if prop.categorizacion is None else prop.categorizacion.codigo,
            'direccion': prop.direccion,
            'referencia': prop.referencia,
            'observaciones': prop.observaciones,
            'habilitado': prop.habilitado,
            'anio': prop.anio,
            'periodo': None if prop.periodo is None else prop.periodo.periodo,
            'area_terreno': float(total_area_terreno2024),
            'area_construccion': float(total_area_construccion2024),
            'tipo': None if prop.tipo is None else prop.tipo.descripcion,
            'propietarios': propietarios  # Concatenated owners
        })

    return JsonResponse(data, safe=False) 


@csrf_exempt
def listar_inmuebles(request):
    today = date.today()  # Año vigente para evaluar si los pagos están vigentes o no
    inicio=datetime.now().strftime("%Y-%m-%d_%H-%M-%S") 

    # Prefetch related data
    inmuebles_filtrados = Inmueble.objects.prefetch_related(
    #inmuebles_filtrados = Inmueble.objects.filter(
    #Q(anio__lt=today.year)).prefetch_related(
    #Q(anio__lte=today.year) | Q(anio__isnull=False)).prefetch_related(
        Prefetch('inmueblevaloracionterreno2024_set', queryset=InmuebleValoracionTerreno2024.objects.all(), to_attr='valoracion_terreno'),
        Prefetch('inmueblepropietarios_set', queryset=InmueblePropietarios.objects.select_related('propietario'), to_attr='propietarios_list')
    ) #[:20]

    data = []
    nTotal=inmuebles_filtrados.count()
    n=0

    for prop in inmuebles_filtrados:
        n=n+1

        baseCalculo2023=0
        TotalImpuesto2023=0
        baseCalculo=0
        TotalImpuesto=0 
        valor_error=0
        mensaje=''

        print('Expediente',prop.numero_expediente,'Total=',nTotal,'Procesando=',n)

        #datosinmueble ={'inmueble': int(prop.numero_expediente)}
        ##valida si el expeidnte esta sano para calcular sus impuestos
        #salud=Datos_Inmuebles_Public(datosinmueble)
        #json_data = salud.data
        #try:
        #    primer_elemento = json_data[0]
        #    mensaje=primer_elemento["mensaje"]
        #    valor_error = primer_elemento["error"]
        #except KeyError:
        #    primer_elemento = json_data
        #    mensaje=primer_elemento["mensaje"]
        #    valor_error = primer_elemento["error"]
        #if valor_error==0: 

        if isinstance(prop.anio, (int, float)):
            if prop.anio<=2023:
                try:
                    deuda2023=Impuesto_Inmueble2023({'inmueble': prop.id, 'anio': prop.anio, 'periodo': prop.periodo.periodo if prop.periodo else 1})
                    json_data_deuda2023 = deuda2023.data
                    #print(json_data_deuda2023)
                    baseCalculo2023 = json_data_deuda2023[0]['cabacera']['basecalculobs']
                    TotalImpuesto2023 = json_data_deuda2023[0]['cabacera']['total']
                    #print('TOTAL',TotalImpuesto2023,baseCalculo2023) 
                except Exception as e:
                    mensaje = 'Error: ' + str(e)
            try:
                deuda=Impuesto_Inmueble({'inmueble': prop.id, 'anio': 2024, 'periodo': prop.periodo.periodo if prop.periodo else 1})
                json_data_deuda = deuda.data
                #print(json_data_deuda)
                baseCalculo = json_data_deuda[0]['cabacera']['basecalculobs']
                TotalImpuesto = json_data_deuda[0]['cabacera']['total']
                #print('TOTAL',TotalImpuesto,baseCalculo) 

            except Exception as e:
                mensaje = 'Error: ' + str(e)
                

            terreno2024 = prop.valoracion_terreno[0] if prop.valoracion_terreno else None
            total_area_terreno2024 = terreno2024.area if terreno2024 and terreno2024.area else 0
            uso_area_terreno2024 = terreno2024.tipologia_categorizacion.descripcion if terreno2024 and terreno2024.tipologia_categorizacion else 0

            total_area_construccion2024 = 0

            if terreno2024:
                construccion2024 = InmuebleValoracionConstruccion2024.objects.filter(inmueblevaloracionterreno=terreno2024)
                total_area_construccion2024 = construccion2024.aggregate(total_area=Sum('area'))['total_area'] or 0
                if total_area_construccion2024:

                    uso_construccion2024 = ", ".join([f"{uso.tipologia_categorizacion.descripcion} ({float(uso.area)})" for uso in construccion2024])
                else:
                    uso_construccion2024='SIN CONSTRUCCION'
        else:
            mensaje='ERROR FALTA AÑO DE ULTIMO PAGO'

        # Obtener y concatenar propietarios
        propietarios = ", ".join([
            f"{propietario.propietario.nombre} ({propietario.propietario.numero_documento})"
            for propietario in prop.propietarios_list
        ])

        # Verificar si el año es None antes de la comparación
        estatus = "Vigente" if prop.anio and prop.anio > today.year else "No Vigente"
        #print(estatus,prop.anio)
        data.append({ 
            'estatus': estatus,
            'expediente': prop.numero_expediente,
            'tipo': prop.tipo.descripcion if prop.tipo else None,
            'zona': prop.zona.codigo if prop.zona else None,
            'categorizacion': prop.categorizacion.codigo if prop.categorizacion else None,
            'direccion': prop.direccion,
            'referencia': prop.referencia,
            'observaciones': prop.observaciones,
            'habilitado': prop.habilitado,
            'anio': prop.anio if isinstance(prop.anio, (int, float)) else None,
            'periodo': prop.periodo.periodo if prop.periodo else None,
            'area_terreno': float(total_area_terreno2024) ,
            'uso_terreno':uso_area_terreno2024,
            'area_construccion': float(total_area_construccion2024),
            'uso_construccion': uso_construccion2024,
            'propietarios': propietarios,  # Concatenated owners
            'mensaje':mensaje,
            'impuesto2023':float(TotalImpuesto2023) if isinstance(prop.anio, (int, float)) else None,
            'BaseCalculo2023':float(baseCalculo2023) if isinstance(prop.anio, (int, float)) else None,
            'impuesto':float(TotalImpuesto) if isinstance(prop.anio, (int, float)) else None,
            'BaseCalculo':float(baseCalculo) if isinstance(prop.anio, (int, float)) else None,
            'TOTAL impuesto':float(TotalImpuesto+TotalImpuesto2023) if isinstance(prop.anio, (int, float)) else None


        })

    print('inicio',inicio)
    print('Fin',datetime.now().strftime("%Y-%m-%d_%H-%M-%S") ) 

    df = pd.DataFrame(data)

    # Escribir el DataFrame en un archivo Excel
    excel_buffer = BytesIO()
    df.to_excel(excel_buffer, index=False)
    excel_buffer.seek(0) 



    # Obtener la fecha y hora actual  
    fecha_hora_actual = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")  # Formato: YYYY-MM-DD_HH-MM-SS  

    # Concatenar la fecha y hora al nombre del archivo  
    nombre_archivo = f'EstadisticaHacienda_{fecha_hora_actual}.xlsx'  

    # Guardar el archivo  
    excel_document = ExcelDocument(title=f'EstadisticaHacienda_{fecha_hora_actual}' , excel_file=nombre_archivo)  
    excel_document.excel_file.save(nombre_archivo, ContentFile(excel_buffer.getvalue()))  

    # Guardar el archivo en el modelo ExcelDocument
    #excel_document = ExcelDocument(title='Mi Excel', excel_file='mi_excel.xlsx')
    #excel_document.excel_file.save('mi_excel.xlsx', ContentFile(excel_buffer.getvalue()))

    print("Archivo Excel guardado en el modelo ExcelDocument.")

    return JsonResponse(data, safe=False) 


@csrf_exempt
def filtrar_flujos(request):
    numero_expediente = request.GET.get('numero_expediente', None)
    flujos_filtrados = Flujo.objects.filter( Q(inmueble__numero_expediente__icontains=numero_expediente))
    data = [
        {'id': prop.id,
         'descripcion_flujo': prop.pagoestadocuenta.liquidacion.tipoflujo.descripcion,
         'expediente': prop.inmueble.numero_expediente,
         'nombre_propietario': prop.pagoestadocuenta.liquidacion.propietario.nombre,
         'estado_display': dict(Flujo.ESTADO)[prop.estado],
         'fecha': prop.fecha
        } 
        for prop in flujos_filtrados
        ]

    return JsonResponse(data, safe=False)

@csrf_exempt
def filtrar_patentes(request):
    numero_expediente = request.GET.get('numero_expediente', None)
    patentes_filtrados = AE_Patente.objects.filter(numero=numero_expediente)

    data = [{'id': prop.id, 
            'numero':prop.numero,
            'tipo_patente':prop.tipo_patente,
            'numero_documento_representante':prop.numero_documento_representante,
            'nombre_representante':prop.nombre_representante,
            'cargo_representante':prop.cargo_representante,
            'telefono':prop.telefono,
            'horario_desde':prop.horario_desde,
            'horario_hasta':prop.horario_hasta,
            'nro_inmuebles':prop.nro_inmuebles,
            'nro_solicitud':prop.nro_solicitud,
            'nro_tomo':prop.nro_tomo,
            'habilitado':prop.habilitado
           } for prop in patentes_filtrados]

    return JsonResponse(data, safe=False)


class TipoInmuebleViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=TipoInmueble.objects.all()
    serializers = {
        'default': TipoInmuebleSerializer
    }


class EstatusInmuebleViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=EstatusInmueble.objects.all()
    serializers = {
        'default': EstatusInmuebleSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'codigo':['exact'],
    }

class NivelInmuebleViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=NivelInmueble.objects.all()
    serializers = {
        'default': NivelInmuebleSerializer
    }

class UnidadInmuebleViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=UnidadInmueble.objects.all()
    serializers = {
        'default': UnidadInmuebleSerializer
    }        

class TipoDocumentoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=TipoDocumento.objects.all()
    serializers = {
        'default': TipoDocumentoSerializer
    }

class TipoEspecialViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=TipoEspecial.objects.all()
    serializers = {
        'default': TipoEspecialSerializer
    }

class TipoTenenciaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=TipoTenencia.objects.all()
    serializers = {
        'default': TipoTenenciaSerializer
    }

class TopografiaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Topografia.objects.all()
    serializers = {
        'default': TopografiaSerializer
    }

class AccesoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Acceso.objects.all()
    serializers = {
        'default': AccesoSerializer
    }

class FormaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Forma.objects.all()
    serializers = {
        'default': FormaSerializer
    }        

class UbicacionViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Ubicacion.objects.all()
    serializers = {
        'default': UbicacionSerializer
    }

class UsoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Uso.objects.all()
    serializers = {
        'default': UsoSerializer
    }

class RegimenViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Regimen.objects.all()
    serializers = {
        'default': RegimenSerializer
    }

class ServiciosViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Servicios.objects.all()
    serializers = {
        'default': ServiciosSerializer
    }

class FinesFiscalesViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=FinesFiscales.objects.all()
    serializers = {
        'default': FinesFiscalesSerializer
    }

class TipoDesincorporacionViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=TipoDesincorporacion.objects.all()
    serializers = {
        'default': TipoDesincorporacionSerializer
    }        

class TipoTransaccionViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=TipoTransaccion.objects.all()
    serializers = {
        'default': TipoTransaccionSerializer
    }

class TipologiaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Tipologia.objects.all()
    serializers = {
        'default': TipologiaSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'zona':['exact'],
    }

class Tipologia_CategorizacionViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Tipologia_Categorizacion.objects.all()
    serializers = {
        'default': Tipologia_CategorizacionSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'categorizacion':['exact'],
    }

class ZonaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Zona.objects.all()
    serializers = {
        'default': ZonaSerializer
    }

class CategorizacionViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Categorizacion.objects.all()
    serializers = {
        'default': CategorizacionSerializer
    }

class InmuebleViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Inmueble.objects.all()
    serializers = {
        'default': InmuebleSerializer
    }


class InmueblePropiedadViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset= InmueblePropiedad.objects.all()
    serializers = {
        'default': InmueblePropiedadSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'inmueble':['exact'],
    }

class InmueblePropietariosViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset= InmueblePropietarios.objects.all()
    serializers = {
        'default': InmueblePropietariosSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'propietario':['exact'],
      'inmueble':['exact'],
    }

class InmuebleTerrenoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset= InmuebleTerreno.objects.all()
    serializers = {
        'default': InmuebleTerrenoSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'inmueble':['exact'],
    }    

class InmuebleTerrenoTopografiaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset= InmuebleTerrenoTopografia.objects.all()
    serializers = {
        'default': InmuebleTerrenoTopografiaSerializer
    }
class InmuebleTerrenoAccesoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset= InmuebleTerrenoAcceso.objects.all()
    serializers = {
        'default': InmuebleTerrenoAccesoSerializer
    }
class InmuebleTerrenoUsoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset= InmuebleTerrenoUso.objects.all()
    serializers = {
        'default': InmuebleTerrenoUsoSerializer
    }

class InmuebleTerrenoServicioViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset= InmuebleTerrenoServicio.objects.all()
    serializers = {
        'default': InmuebleTerrenoServicioSerializer
    }

class UsoConstruccionViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=UsoConstruccion.objects.all()
    serializers = {
        'default': UsoConstruccionSerializer
    }

class SoporteViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Soporte.objects.all()
    serializers = {
        'default': SoporteSerializer
    }

class TechoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Techo.objects.all()
    serializers = {
        'default': TechoSerializer
    }

class CubiertaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Cubierta.objects.all()
    serializers = {
        'default': CubiertaSerializer
    }

class TipoParedViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=TipoPared.objects.all()
    serializers = {
        'default': TipoParedSerializer
    }

class AcabadoParedViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=AcabadoPared.objects.all()
    serializers = {
        'default': AcabadoParedSerializer
    }

class ConservacionViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Conservacion.objects.all()
    serializers = {
        'default': ConservacionSerializer
    }

class InmuebleConstruccionViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset= InmuebleConstruccion.objects.all()
    serializers = {
        'default': InmuebleConstruccionSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'inmueble':['exact'],
    }    

class InmuebleConstruccionSoporteViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=InmuebleConstruccionSoporte.objects.all()
    serializers = {
        'default': InmuebleConstruccionSoporteSerializer
    }

class InmuebleConstruccionTechoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=InmuebleConstruccionTecho.objects.all()
    serializers = {
        'default': InmuebleConstruccionTechoSerializer
    }

class InmuebleConstruccionCubiertaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=InmuebleConstruccionCubierta.objects.all()
    serializers = {
        'default': InmuebleConstruccionCubiertaSerializer
    }

class InmuebleValoracionTerrenoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset= InmuebleValoracionTerreno.objects.all()
    serializers = {
        'default': InmuebleValoracionTerrenoSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'inmueble':['exact'],
    }

class InmuebleValoracionConstruccionViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset= InmuebleValoracionConstruccion.objects.all()
    serializers = {
        'default': InmuebleValoracionConstruccionSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'inmueblevaloracionterreno':['exact'],
    }

class InmuebleValoracionTerreno2024Viewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset= InmuebleValoracionTerreno2024.objects.all()
    serializers = {
        'default': InmuebleValoracionTerreno2024Serializer
     }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'inmueble':['exact'],
    }

class InmuebleValoracionConstruccion2024Viewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset= InmuebleValoracionConstruccion2024.objects.all()
    serializers = {
        'default': InmuebleValoracionConstruccion2024Serializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'inmueblevaloracionterreno':['exact'],
    }


class InmuebleUbicacionViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset= InmuebleUbicacion.objects.all()
    serializers = {
        'default': InmuebleUbicacionSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'inmueble':['exact'],
    }    

class InmuebleFaltanteViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset= InmuebleFaltante.objects.all()
    serializers = {
        'default': InmuebleFaltanteSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'inmueble':['exact'],
    }    

class PropietarioViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset= Propietario.objects.all()
    serializers = {
        'default': PropietarioSerializer
    }



class TasaBCVViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=TasaBCV.objects.all()
    serializers = {
        'default': TasaBCVSerializer
    }

    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'habilitado':['exact'],
    }

class UnidadTributariaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=UnidadTributaria.objects.all()
    serializers = {
        'default': UnidadTributariaSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'habilitado':['exact'],
    }

class MonedaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Moneda.objects.all()
    serializers = {
        'default': MonedaSerializer
    }

class TasaMultaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=TasaMulta.objects.all()
    serializers = {
        'default': TasaMultaSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'aplica':['exact'],
    }

class TipoFlujoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=TipoFlujo.objects.all()
    serializers = {
        'default': TipoFlujoSerializer
    }

class TipoFlujoDetalleViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=TipoFlujoDetalle.objects.all()
    serializers = {
        'default': TipoFlujoDetalleSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'tipoflujo':['exact'],
    }

class EstadoCuentaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=EstadoCuenta.objects.all()
    serializers = {
        'default': EstadoCuentaSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'propietario':['exact'],
      'habilitado':['exact'],
    }

class EstadoCuentaDetalleViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=EstadoCuentaDetalle.objects.all()
    serializers = {
        'default': EstadoCuentaDetalleSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'estadocuenta_id':['exact'],
    }

class LiquidacionViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Liquidacion.objects.all()
    serializers = {
        'default': LiquidacionSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'propietario':['exact'],
      'habilitado':['exact'],
    }

class LiquidacionDetalleViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=LiquidacionDetalle.objects.all()
    serializers = {
        'default': LiquidacionDetalleSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'liquidacion_id':['exact'],
    }

class TipoPagoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=TipoPago.objects.all()
    serializers = {
        'default': TipoPagoSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'codigo':['exact'],
      'lstar':['exact'],
      'codigo':['exact'],
    }

class BancoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Banco.objects.all()
    serializers = {
        'default': BancoSerializer
    }

class BancoCuentaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=BancoCuenta.objects.all()
    serializers = {
        'default': BancoCuentaSerializer
    }


class PagoEstadoCuentaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=PagoEstadoCuenta.objects.all()
    serializers = {
        'default': PagoEstadoCuentaSerializer
    }

class PagoEstadoCuentaDetalleViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=PagoEstadoCuentaDetalle.objects.all()
    serializers = {
        'default': PagoEstadoCuentaDetalleSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'pagoestadocuenta_id':['exact'],
      'tipopago_id':['exact'],
      'bancocuenta':['exact'],
    }
    def get_queryset(self):
        queryset = super().get_queryset()
        bancocuenta_isnull = self.request.query_params.get('bancocuenta__isnull', None)

        if bancocuenta_isnull == 'true':
            queryset = queryset.filter(bancocuenta__isnull=True)
        elif bancocuenta_isnull == 'false':
            queryset = queryset.filter(bancocuenta__isnull=False)

        return queryset


class CorrelativoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Correlativo.objects.all()
    serializers = {
        'default': CorrelativoSerializer
    }

class FlujoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Flujo.objects.all()
    serializers = {
        'default': FlujoSerializer
    }

class FlujoDetalleViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=FlujoDetalle.objects.all()
    serializers = {
        'default': FlujoDetalleSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'estado':['exact'],
      'recibe_usuario':['exact'],
      'departamento_recibe':['exact'],
      'flujo':['exact'],
      'tarea':['exact'],
    }

class IC_PeriodoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=IC_Periodo.objects.all()
    serializers = {
        'default': IC_PeriodoSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'aplica':['exact'],
    }

class IC_ImpuestoCargosViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=IC_ImpuestoCargos.objects.all()
    serializers = {
        'default': IC_ImpuestoCargosSerializer
    }
    
    
class IC_ImpuestoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=IC_Impuesto.objects.all()
    serializers = {
        'default':IC_ImpuestoSerializer
    }


class IC_ImpuestoDetalleViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=IC_ImpuestoDetalle.objects.all()
    serializers = {
        'default': IC_ImpuestoDetalleSerializer
    }


class IC_ImpuestoCorreccionesViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=IC_ImpuestoCorrecciones.objects.all()
    serializers = {
        'default': IC_ImpuestoCorreccionesSerializer
    }

class IC_ImpuestoCorreccionesDetalleViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=IC_ImpuestoCorreccionesDetalle.objects.all()
    serializers = {
        'default': IC_ImpuestoCorreccionesDetalleSerializer
    }


class IC_ImpuestoPeriodoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=IC_ImpuestoPeriodo.objects.all()
    serializers = {
        'default': IC_ImpuestoPeriodoSerializer
    }

class IC_ImpuestoDescuentoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=IC_ImpuestoDescuento.objects.all()
    serializers = {
        'default': IC_ImpuestoDescuentoSerializer
    }

class IC_ImpuestoExoneracionViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=IC_ImpuestoExoneracion.objects.all()
    serializers = {
        'default': IC_ImpuestoExoneracionSerializer
    }

class IC_ImpuestoDetalleDescuentosViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=IC_ImpuestoDetalleDescuentos.objects.all()
    serializers = {
        'default': IC_ImpuestoDetalleDescuentosSerializer
    }
class IC_ImpuestoDetalleMoraViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=IC_ImpuestoDetalleMora.objects.all()
    serializers = {
        'default': IC_ImpuestoDetalleMoraSerializer
    }

class AE_ActividadEconomicaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=AE_ActividadEconomica.objects.all()
    serializers = {
        'default': AE_ActividadEconomicaSerializer
    }

class AE_ActividadEconomicaDetalleViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=AE_ActividadEconomicaDetalle.objects.all()
    serializers = {
        'default': AE_ActividadEconomicaDetalleSerializer
    }


class AE_PatenteViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=AE_Patente.objects.all()
    serializers = {
        'default': AE_PatenteSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'propietario':['exact'],
    }

class AE_Patente_ActividadEconomicaViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=AE_Patente_ActividadEconomica.objects.all()
    serializers = {
        'default': AE_Patente_ActividadEconomicaSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'AE_patente':['exact'],
      'AE_actividadeconomica':['exact'],
    }

class TasaInteresViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=TasaInteres.objects.all()
    serializers = {
        'default': TasaInteresSerializer
    }

#class NotaCreditoViewset(MultiSerializerViewSet):
#    permission_classes = [IsAuthenticated]
#    queryset=NotaCredito.objects.all()
#    serializers = {
#        'default': NotaCreditoSerializer
#    }
#    filter_backends = [DjangoFilterBackend]
#    filterset_fields = {
#      'propietario':['exact'],
#      'saldo':['exact'],
#    }


class NotaCreditoFilter(filters.FilterSet):
    saldo_gt = filters.NumberFilter(field_name='saldo', lookup_expr='gt')

    class Meta:
        model = NotaCredito
        fields = ['propietario', 'saldo_gt','pagoestadocuenta']

        

class NotaCreditoViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset = NotaCredito.objects.all()
    serializers = {
        'default': NotaCreditoSerializer
    }
    filter_backends = [DjangoFilterBackend]
    filterset_class = NotaCreditoFilter



class ExcelDocumentViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=ExcelDocument.objects.all()
    serializers = {
        'default': ExcelDocumentSerializer
    }

class ExcelDocumentLOGViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=ExcelDocumentLOG.objects.all()
    serializers = {
        'default': ExcelDocumentLOGSerializer
    }


class CorridasBancariasViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=CorridasBancarias.objects.all()
    serializers = {
        'default': CorridasBancariasSerializer
    }
    
    filter_backends = [DjangoFilterBackend]
    filterset_fields = {
      'bancocuenta':['exact'],
      'situado':['exact'],
      'habilitado':['exact'],
    }
class ComunidadViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=Comunidad.objects.all()
    serializers = {
        'default': ComunidadSerializer
    }

class InmuebleCategorizacionViewset(MultiSerializerViewSet):
    permission_classes = [IsAuthenticated]
    queryset=InmuebleCategorizacion.objects.all()
    serializers = {
        'default': InmuebleCategorizacionSerializer
    }

class CorridasBancariaSinPagoView(APIView):
    def get(self, request):
        corridaancaria_sin_pago = CorridasBancarias.objects.filter(habilitado=True).exclude(
            Q(id__in=PagoEstadoCuentaDetalle.objects.values_list('corridasbancarias_id', flat=True).exclude(corridasbancarias_id__isnull=True))
        )
        serializer = CorridasBancariasSerializer(corridaancaria_sin_pago, many=True)
        return Response(serializer.data)

class CorridasBancariaSinPagoRecaudosView(APIView):
    def get(self, request):
        corridaancaria_sin_pago = CorridasBancarias.objects.exclude(
            Q(id__in=PagoEstadoCuentaDetalle.objects.values_list('corridasbancarias_id', flat=True).exclude(corridasbancarias_id__isnull=True))
        )
        serializer = CorridasBancariasSerializer(corridaancaria_sin_pago, many=True)
        return Response(serializer.data)


#class Viewset(MultiSerializerViewSet):
#    permission_classes = [IsAuthenticated]
#    queryset=.objects.all()
#    serializers = {
#        'default': Serializer
#    }